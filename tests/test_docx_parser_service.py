from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import docx

from app.core.config import settings
from app.services.docx_parser_service import parse_docx


def create_sample_docx(
    file_path: Path,
    paragraphs: list[str],
    table_data: list[list[str]] | None = None,
) -> Path:
    """Helper function to create a test DOCX with paragraphs and optional table data."""
    doc = docx.Document()
    for p in paragraphs:
        if p:
            doc.add_paragraph(p)

    if table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                table.cell(r_idx, c_idx).text = val

    doc.save(str(file_path))
    return file_path


class TestDOCXParserService(unittest.TestCase):
    """Test suite for DOCX Parser Service."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.tmp_path = Path(self.temp_dir.name)
        self.processed_dir = self.tmp_path / "processed"

        # Patch settings.processed_directory for isolated testing
        self.patcher = patch.object(settings, "processed_directory", self.processed_dir)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()
        self.temp_dir.cleanup()

    def test_parse_docx_valid(self):
        """Test successful parsing of a valid DOCX document with text and table data."""
        upload_dir = self.tmp_path / "uploads"
        upload_dir.mkdir()

        docx_file = upload_dir / "sample_doc.docx"
        create_sample_docx(
            docx_file,
            paragraphs=["Hello world from paragraph 1.", "Welcome to paragraph 2."],
            table_data=[["Header 1", "Header 2"], ["Value 1", "Value 2"]],
        )

        result_path = parse_docx(docx_file)

        self.assertEqual(result_path, self.processed_dir / "sample_doc.txt")
        self.assertTrue(result_path.exists())

        content = result_path.read_text(encoding="utf-8")
        self.assertIn("Hello world from paragraph 1.", content)
        self.assertIn("Welcome to paragraph 2.", content)
        self.assertIn("Header 1 | Header 2", content)
        self.assertIn("Value 1 | Value 2", content)

    def test_parse_docx_file_not_found(self):
        """Test that FileNotFoundError is raised when input DOCX file does not exist."""
        non_existent = self.tmp_path / "non_existent.docx"
        with self.assertRaises(FileNotFoundError):
            parse_docx(non_existent)

    def test_parse_docx_no_extractable_text(self):
        """Test that ValueError is raised when DOCX contains no extractable text."""
        upload_dir = self.tmp_path / "uploads"
        upload_dir.mkdir()

        docx_file = upload_dir / "empty.docx"
        create_sample_docx(docx_file, ["", "   "])

        with self.assertRaises(ValueError):
            parse_docx(docx_file)

    def test_parse_docx_corrupt_file_propagates_exception(self):
        """Test that unexpected parsing errors (e.g. corrupt file) propagate."""
        corrupt_file = self.tmp_path / "corrupt.docx"
        corrupt_file.write_bytes(b"This is not a valid DOCX zip package.")

        with self.assertRaises(Exception):
            parse_docx(corrupt_file)


if __name__ == "__main__":
    unittest.main()
