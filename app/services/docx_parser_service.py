"""DOCX parser service module for extracting text from DOCX documents."""

from pathlib import Path
import docx

from app.core.config import settings


def parse_docx(docx_path: Path) -> Path:
    """Parse text content from a DOCX file and save it as a plain text file.

    Extracts paragraph text and table content and writes the combined output
    to a text file in the processed directory using the same stem as the input DOCX.

    Args:
        docx_path: Path to the input DOCX file.

    Returns:
        Path: Path to the generated .txt file in the processed directory.

    Raises:
        FileNotFoundError: If the docx_path does not exist or is not a file.
        ValueError: If the DOCX contains no extractable text.
        Exception: Propagates any unexpected python-docx or file system errors.
    """
    if not docx_path.exists() or not docx_path.is_file():
        raise FileNotFoundError(f"DOCX file does not exist: {docx_path}")

    target_dir = settings.processed_directory
    extracted_blocks: list[str] = []

    doc = docx.Document(str(docx_path))

    # Extract text from paragraphs
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            extracted_blocks.append(text)

    # Extract text from tables
    for table in doc.tables:
        table_rows_text: list[str] = []
        for row in table.rows:
            row_cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells_text:
                table_rows_text.append(" | ".join(row_cells_text))
        if table_rows_text:
            extracted_blocks.append("\n".join(table_rows_text))

    if not extracted_blocks:
        raise ValueError(f"DOCX file contains no extractable text: {docx_path}")

    combined_text = "\n\n".join(extracted_blocks)

    target_dir.mkdir(parents=True, exist_ok=True)
    output_file_path = target_dir / f"{docx_path.stem}.txt"
    output_file_path.write_text(combined_text, encoding="utf-8")

    return output_file_path
